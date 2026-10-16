import fitz
import docx
import pandas as pd
import pytest

from assumption_monitor import reader

HEADER = "ACME Corp | Business Plan 2026 | CONFIDENTIAL"
FOOTER = "ACME Business Plan - Draft v3"


def page(n, body):
    """One fake PDF page: header, footer, page label, then the body lines."""
    return f"{HEADER}\n{FOOTER}\nPage {n}\n{body}"


def run(pages, doc_id="t"):
    """Clean fake pages (one string per page) without opening any file."""
    lines = []
    for n, text in enumerate(pages, start=1):
        for raw in text.split("\n"):
            lines.append(reader.Line(raw, n))
    return reader.clean_lines(lines, len(pages), doc_id)


# ---------- Acceptance checks from section 4 ----------
def test_headers_footers_and_page_labels_are_removed():
    pages = [
        page(1, "1. Overview\nWe assume diesel prices will remain stable across\nall regions this year."),
        page(2, "2.1 Fuel costs\nWe expect demand to grow by 8% next year."),
        page(3, "Management plans to open ten new stores."),
    ]
    sentences, audit = run(pages)
    text = " ".join(sentences["sentence"])
    assert "CONFIDENTIAL" not in text
    assert "Draft v3" not in text
    assert "Page 2" not in text
    assert {"header_footer", "page_label"} <= set(audit["reason"])


def test_rs_crore_and_abbreviations_stay_intact():
    body = "In 2025 the company earned revenue of Rs. 42.5 crore and a margin of 9.2%. ACME Pvt. Ltd. plans to grow."
    sentences, _ = run([page(1, body), page(2, "Filler line for page two."), page(3, "Filler line for page three.")])
    first = sentences["sentence"].iloc[0]
    assert first == "In 2025 the company earned revenue of Rs. 42.5 crore and a margin of 9.2%."
    assert "ACME Pvt. Ltd. plans to grow." in sentences["sentence"].tolist()
    assert sentences["money"].iloc[0] == ["Rs. 42.5 crore"]
    assert sentences["percents"].iloc[0] == ["9.2%"]


def test_headings_become_sections_not_sentences():
    pages = [
        page(1, "1. Overview\nWe assume diesel prices will remain stable this year."),
        page(2, "2.1 Fuel costs\nWe expect demand to grow by 8% next year."),
        page(3, "2.2 Labor\nWages are expected to rise by 6% per year."),
    ]
    sentences, _ = run(pages)
    by_section = dict(zip(sentences["section"], sentences["sentence"]))
    assert "2.1 Fuel costs" in by_section
    assert "2.2 Labor" in by_section
    assert "2.1 Fuel costs" not in sentences["sentence"].tolist()


def test_page_references_survive_cleaning_and_page_spans_work():
    pages = [
        page(1, "We assume diesel prices will remain stable this year."),
        page(2, "2.1 Currency\nWe expect the exchange rate to remain close to"),
        page(3, "INR 83 per USD during the plan period. Management plans to open ten new stores."),
    ]
    sentences, _ = run(pages)
    rows = sentences.set_index("sentence")
    assert rows.loc["We assume diesel prices will remain stable this year.", "page_start"] == 1
    crossing = "We expect the exchange rate to remain close to INR 83 per USD during the plan period."
    assert rows.loc[crossing, "page_start"] == 2
    assert rows.loc[crossing, "page_end"] == 3
    assert rows.loc["Management plans to open ten new stores.", "page_start"] == 3


def test_isolated_short_lines_are_kept():
    body = "The company plans\nto\nexpand its fleet next year."
    sentences, audit = run([page(1, body), page(2, "Filler line for page two."), page(3, "Filler line for page three.")])
    assert "The company plans to expand its fleet next year." in sentences["sentence"].tolist()
    assert "table_fragment" not in set(audit["reason"])


def test_table_cells_are_excluded_but_logged():
    body = "Revenue is shown below.\nItem\n2026\n2027\nRevenue\n52.0\nWe assume freight rates will rise by 5% per year."
    sentences, audit = run([page(1, body), page(2, "Filler line for page two."), page(3, "Filler line for page three.")])
    joined = " ".join(sentences["sentence"])
    assert "2027" not in joined
    logged = audit[audit["reason"] == "table_fragment"]["text"].tolist()
    assert "2026" in logged and "52.0" in logged
    assert "We assume freight rates will rise by 5% per year." in sentences["sentence"].tolist()


def test_duplicates_are_merged_and_keep_all_page_references():
    pages = [
        page(1, "We assume diesel prices will remain stable this year."),
        page(2, "Management plans to open ten new stores."),
        page(3, "We assume diesel prices will remain stable this year."),
    ]
    sentences, _ = run(pages)
    dup = sentences[sentences["sentence"].str.startswith("We assume diesel")]
    assert len(dup) == 1
    assert dup["pages"].iloc[0] == "1,3"
    assert dup["n_occurrences"].iloc[0] == 2


def test_repeated_real_sentences_are_protected_from_header_removal():
    keep = "We assume demand will hold."
    pages = [page(n, f"{keep}\nOn page {n} the company reviews its plans for the year.") for n in (1, 2, 3)]
    sentences, audit = run(pages)
    assert keep in sentences["sentence"].tolist()
    assert keep not in audit["text"].tolist()


def test_short_sentences_are_flagged_not_deleted():
    sentences, _ = run([page(1, "Revenue grew."), page(2, "Filler line for page two."), page(3, "Filler line for page three.")])
    row = sentences[sentences["sentence"] == "Revenue grew."].iloc[0]
    assert row["extraction_status"] == "too_short"


def test_sentence_record_has_required_fields_and_stable_ids():
    sentences, _ = run([page(1, "We assume diesel prices will remain stable this year."), page(2, "Filler line for page two."), page(3, "Filler line for page three.")], doc_id="plan9")
    for col in ["sentence_id", "doc_id", "page_start", "page_end", "section", "sentence",
                "n_words", "percents", "money", "extraction_status"]:
        assert col in sentences.columns
    assert sentences["sentence_id"].iloc[0] == "plan9-0001"


# ---------- Real files: PDF ----------
def make_text_pdf(path, n_pages=3):
    doc = fitz.open()
    for n in range(1, n_pages + 1):
        p = doc.new_page()
        p.insert_text((50, 50), HEADER, fontsize=9)
        p.insert_text((50, 100), f"We expect revenue to grow by {n}0% next year.", fontsize=11)
        p.insert_text((50, 800), f"Page {n}", fontsize=9)
    doc.save(str(path))
    doc.close()


def test_pdf_end_to_end(tmp_path):
    path = tmp_path / "sample.pdf"
    make_text_pdf(path)
    result = reader.read_document(path)
    text = " ".join(result.sentences["sentence"])
    assert "CONFIDENTIAL" not in text
    assert result.sentences["page_start"].tolist() == [1, 2, 3]
    assert result.sentences["doc_id"].iloc[0] == "sample"


def test_scanned_pdf_gives_a_clear_message(tmp_path):
    path = tmp_path / "scan.pdf"
    doc = fitz.open()
    doc.new_page()
    doc.save(str(path))
    doc.close()
    with pytest.raises(reader.UnsupportedFileError) as err:
        reader.read_document(path)
    assert "scanned" in str(err.value).lower()


def test_broken_pdf_gives_a_clear_message(tmp_path):
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"this is not a pdf")
    with pytest.raises(reader.ExtractionError):
        reader.read_document(path)


# ---------- Real files: Word ----------
def test_docx_end_to_end(tmp_path):
    path = tmp_path / "plan.docx"
    d = docx.Document()
    d.add_heading("Fuel costs", level=1)
    d.add_paragraph("We assume diesel prices will remain stable. We expect demand to grow by 8% next year.")
    d.add_table(rows=2, cols=2)
    d.save(str(path))

    result = reader.read_document(path)
    assert len(result.sentences) == 2
    assert set(result.sentences["section"]) == {"Fuel costs"}
    assert pd.isna(result.sentences["page_start"].iloc[0])
    assert any("table" in w for w in result.warnings)


# ---------- Errors and command line ----------
def test_unsupported_and_missing_files(tmp_path):
    txt = tmp_path / "a.txt"
    txt.write_text("hello")
    with pytest.raises(reader.UnsupportedFileError):
        reader.read_document(txt)
    with pytest.raises(reader.ExtractionError):
        reader.read_document(tmp_path / "missing.pdf")


def test_command_line_reports_errors_without_crashing(tmp_path, capsys):
    txt = tmp_path / "a.txt"
    txt.write_text("hello")
    assert reader.main([str(txt), "--out-dir", str(tmp_path)]) == 1
    assert "Unsupported" in capsys.readouterr().err
