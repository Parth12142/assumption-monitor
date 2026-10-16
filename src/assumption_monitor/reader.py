"""Read PDF and Word plans into a clean sentence table (implementation plan, section 4)."""
from __future__ import annotations

import bisect
import re
import sys
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd


# ---------- Errors ----------
class ExtractionError(Exception):
    """The file could not be found, opened or read."""


class UnsupportedFileError(Exception):
    """The file type or content is not supported (for example a scanned PDF)."""


# ---------- Data containers ----------
@dataclass
class Line:
    text: str
    page: int | None = None   # None for Word files (no fixed pages)
    heading: bool = False     # True when the source marks it as a heading (Word styles)
    para_end: bool = False    # True when the line ends a paragraph in the source (Word)


@dataclass
class ReadResult:
    sentences: pd.DataFrame
    audit: pd.DataFrame
    warnings: list = field(default_factory=list)


# ---------- Settings ----------
PAGE_LABEL = re.compile(r"^(page\s+\d+(\s+of\s+\d+)?|-\s*\d+\s*-)$", re.IGNORECASE)
HEADING = re.compile(r"^\d+(\.\d+)*\.?\s+[A-Z]")
SENTENCE_END = (".", "!", "?")
SHORT_LINE = 25      # characters
MIN_TABLE_RUN = 3    # this many short lines in a row look like table cells
MIN_WORDS = 5

ABBREVIATIONS = ["Rs.", "Pvt.", "Ltd.", "No.", "Dr.", "Mr.", "Mrs.", "vs.", "approx."]
DOT = "\u2024"       # a look-alike dot, same length as ".", used to protect abbreviations
SPLIT = re.compile(r"(?<=[.!?])\s+(?=[A-Z])")

PERCENT = re.compile(r"\d+(?:\.\d+)?\s?%")
MONEY = re.compile(
    r"(?:Rs\.|USD|INR|\$)\s?\d[\d,]*(?:\.\d+)?(?:\s?(?:crore|lakh|million|billion))?"
)

COLUMNS = [
    "sentence_id", "doc_id", "page_start", "page_end", "pages", "section",
    "sentence", "n_words", "n_occurrences", "percents", "money", "extraction_status",
]


# ---------- Step 1: read the file ----------
def load_pdf(path):
    import pymupdf  # PyMuPDF

    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open PDF '{path}': {exc}") from exc

    n_pages = len(doc)
    if n_pages == 0:
        raise ExtractionError(f"PDF '{path}' has no pages.")

    lines, thin_pages = [], []
    for page_no, page in enumerate(doc, start=1):
        text = page.get_text()
        if len(text.strip()) < 20:
            thin_pages.append(page_no)
        for raw in text.split("\n"):
            lines.append(Line(raw, page_no))

    if len(thin_pages) == n_pages:
        raise UnsupportedFileError(
            "No extractable text found. This looks like a scanned or image-only PDF; "
            "OCR is not supported yet."
        )
    warnings = []
    if thin_pages:
        warnings.append(f"Pages with little or no text (possibly scanned): {thin_pages}")
    return lines, n_pages, warnings


def load_docx(path):
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open Word file '{path}': {exc}") from exc

    lines = []
    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name if para.style is not None else "") or ""
        is_heading = style.lower().startswith("heading") or style.lower() == "title"
        lines.append(Line(text, page=None, heading=is_heading, para_end=True))

    if not lines:
        raise UnsupportedFileError("The Word file has no readable paragraphs.")

    warnings = ["Word files have no fixed pages; page numbers are empty."]
    if document.tables:
        warnings.append(f"{len(document.tables)} table(s) were skipped (tables are not read yet).")
    return lines, 0, warnings


# ---------- Step 2: clean the lines ----------
def find_repeated_lines(lines, n_pages, min_share=0.6):
    """Lines that appear on at least 60% of pages are probably headers or footers.

    Long lines, and lines that end like a sentence, are protected: they may be real content.
    """
    if n_pages < 3:
        return set()
    pages_seen = {}
    for ln in lines:
        text = ln.text.strip()
        if text and ln.page is not None:
            pages_seen.setdefault(text, set()).add(ln.page)
    limit = min_share * n_pages
    return {
        text for text, pages in pages_seen.items()
        if len(pages) >= limit and len(text) <= 120 and not text.endswith(SENTENCE_END)
    }


def _is_heading(ln):
    return ln.heading or (
        bool(HEADING.match(ln.text)) and not ln.text.endswith(SENTENCE_END) and len(ln.text) <= 80
    )


def _is_short(ln):
    return (not ln.para_end and not _is_heading(ln) and len(ln.text) <= SHORT_LINE
            and not ln.text.endswith(SENTENCE_END))


def _build_blocks(lines, n_pages, doc_id, audit):
    repeated = find_repeated_lines(lines, n_pages)

    kept = []
    for ln in lines:
        text = ln.text.strip()
        if not text:
            continue
        if text in repeated:
            audit.append({"doc_id": doc_id, "page": ln.page, "text": text, "reason": "header_footer"})
        elif PAGE_LABEL.match(text):
            audit.append({"doc_id": doc_id, "page": ln.page, "text": text, "reason": "page_label"})
        else:
            kept.append(Line(text, ln.page, ln.heading, ln.para_end))

    # Table cells arrive as several short lines in a row. One isolated short line is kept.
    drop, i = set(), 0
    while i < len(kept):
        if _is_short(kept[i]):
            j = i
            while j < len(kept) and _is_short(kept[j]):
                j += 1
            if j - i >= MIN_TABLE_RUN:
                drop.update(range(i, j))
            i = j
        else:
            i += 1

    blocks, current = [], []
    for idx, ln in enumerate(kept):
        if idx in drop:
            audit.append({"doc_id": doc_id, "page": ln.page, "text": ln.text, "reason": "table_fragment"})
            continue
        if _is_heading(ln):
            if current:
                blocks.append(("text", current))
                current = []
            blocks.append(("heading", [ln]))
            continue
        current.append(ln)
        if ln.para_end:
            blocks.append(("text", current))
            current = []
    if current:
        blocks.append(("text", current))
    return blocks


# ---------- Step 3: split into sentences ----------
def _protect(text):
    for abbr in ABBREVIATIONS:
        text = re.sub(
            rf"(?<![A-Za-z]){re.escape(abbr)}",
            lambda m: m.group().replace(".", DOT),
            text,
        )
    return text


def split_with_offsets(text):
    """Return (sentence, start, end) for each sentence, offsets pointing into `text`."""
    safe = _protect(text)
    spans, prev = [], 0
    for m in SPLIT.finditer(safe):
        spans.append((prev, m.start()))
        prev = m.end()
    spans.append((prev, len(safe)))

    result = []
    for start, end in spans:
        chunk = safe[start:end]
        stripped = chunk.strip()
        if stripped:
            first = start + (len(chunk) - len(chunk.lstrip()))
            result.append((stripped.replace(DOT, "."), first, first + len(stripped) - 1))
    return result


def _page_lookup(block_lines):
    text, starts, pages = "", [], []
    for ln in block_lines:
        if text:
            text += " "
        starts.append(len(text))
        pages.append(ln.page)
        text += ln.text
    return text, starts, pages


# ---------- Step 4: put it all in one table ----------
def clean_lines(lines, n_pages, doc_id):
    audit, rows = [], []
    section = "(start)"
    for kind, block in _build_blocks(lines, n_pages, doc_id, audit):
        if kind == "heading":
            section = block[0].text
            continue
        text, starts, pages = _page_lookup(block)
        for sentence, first, last in split_with_offsets(text):
            rows.append({
                "doc_id": doc_id,
                "page_start": pages[bisect.bisect_right(starts, first) - 1],
                "page_end": pages[bisect.bisect_right(starts, last) - 1],
                "section": section,
                "sentence": sentence,
            })

    sentences = _dedupe_and_enrich(rows, doc_id)
    audit_df = pd.DataFrame(audit, columns=["doc_id", "page", "text", "reason"])
    return sentences, audit_df


def _dedupe_and_enrich(rows, doc_id):
    kept, pages_by_key, count_by_key = [], {}, {}
    for row in rows:
        key = " ".join(row["sentence"].lower().split())
        refs = pages_by_key.setdefault(key, set())
        for p in (row["page_start"], row["page_end"]):
            if p is not None:
                refs.add(p)
        count_by_key[key] = count_by_key.get(key, 0) + 1
        if count_by_key[key] == 1:
            kept.append((key, row))

    out = []
    for n, (key, row) in enumerate(kept, start=1):
        n_words = len(row["sentence"].split())
        out.append({
            "sentence_id": f"{doc_id}-{n:04d}",
            "doc_id": doc_id,
            "page_start": row["page_start"],
            "page_end": row["page_end"],
            "pages": ",".join(str(p) for p in sorted(pages_by_key[key])),
            "section": row["section"],
            "sentence": row["sentence"],
            "n_words": n_words,
            "n_occurrences": count_by_key[key],
            "percents": PERCENT.findall(row["sentence"]),
            "money": MONEY.findall(row["sentence"]),
            "extraction_status": "ok" if n_words >= MIN_WORDS else "too_short",
        })
    df = pd.DataFrame(out, columns=COLUMNS)
    df["page_start"] = df["page_start"].astype("Int64")
    df["page_end"] = df["page_end"].astype("Int64")
    return df


# ---------- Public functions ----------
def read_document(path, doc_id=None):
    path = Path(path)
    if not path.exists():
        raise ExtractionError(f"File not found: {path}")
    suffix = path.suffix.lower()
    doc_id = doc_id or path.stem

    if suffix == ".pdf":
        lines, n_pages, warnings = load_pdf(path)
    elif suffix == ".docx":
        lines, n_pages, warnings = load_docx(path)
    else:
        raise UnsupportedFileError(f"Unsupported file type '{suffix}'. Use .pdf or .docx.")

    sentences, audit = clean_lines(lines, n_pages, doc_id)
    if sentences.empty:
        raise ExtractionError("No sentences were found after cleaning.")
    return ReadResult(sentences, audit, warnings)


def preview(result, n_chars=800):
    """First part of the cleaned text, for a quick human check of reader quality."""
    ok = result.sentences[result.sentences["extraction_status"] == "ok"]
    return " ".join(ok["sentence"])[:n_chars]


def main(argv=None):
    import argparse

    parser = argparse.ArgumentParser(description="Turn a PDF or DOCX plan into a sentence table.")
    parser.add_argument("file")
    parser.add_argument("--out-dir", default="data/clean")
    parser.add_argument("--preview", action="store_true")
    args = parser.parse_args(argv)

    try:
        result = read_document(args.file)
    except (ExtractionError, UnsupportedFileError) as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 1

    doc_id = result.sentences["doc_id"].iloc[0]
    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    result.sentences.to_csv(out_dir / f"{doc_id}_sentences.csv", index=False)
    result.audit.to_csv(out_dir / f"{doc_id}_audit.csv", index=False)

    ok = int((result.sentences["extraction_status"] == "ok").sum())
    print(f"{len(result.sentences)} sentences ({ok} ok) saved to {out_dir}/{doc_id}_sentences.csv")
    print(f"{len(result.audit)} removed lines logged in {out_dir}/{doc_id}_audit.csv")
    print(result.audit["reason"].value_counts().to_string() if len(result.audit) else "nothing removed")
    for w in result.warnings:
        print("WARNING:", w)
    if args.preview:
        print("\nPREVIEW:\n" + preview(result))
    return 0


if __name__ == "__main__":
    sys.exit(main())
